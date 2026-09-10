"""
document_builder.py
Modelin JSON çıktısını veya düz metni profesyonel bir .docx dosyasına dönüştürür.
ProjeDanışmanAI v2 — Qwen3-14B
"""

import io
import json
import re
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ── Renkler ──────────────────────────────────────────────────────────────────
KOYU_MAVI  = RGBColor(0x1F, 0x38, 0x64)
ORTA_MAVI  = RGBColor(0x2E, 0x55, 0x97)
GRI        = RGBColor(0x66, 0x66, 0x66)


def _set_color(run, rgb: RGBColor):
    run.font.color.rgb = rgb


def _add_alt_cizgi(paragraf):
    """Paragrafa alt çizgi (border) ekle."""
    pPr = paragraf._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E5597")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _sayfa_ayarlari(doc: Document):
    section = doc.sections[0]
    section.page_width   = Cm(21)
    section.page_height  = Cm(29.7)
    section.left_margin  = Cm(3)
    section.right_margin = Cm(2.5)
    section.top_margin   = Cm(2.5)
    section.bottom_margin = Cm(2.5)


def _baslik_ekle(doc: Document, metin: str, seviye: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(6)
    run = p.add_run(metin)
    run.font.name = "Arial"
    run.font.bold = True
    if seviye == 1:
        run.font.size = Pt(14)
        _set_color(run, KOYU_MAVI)
        _add_alt_cizgi(p)
    else:
        run.font.size = Pt(12)
        _set_color(run, ORTA_MAVI)


def _paragraf_ekle(doc: Document, metin: str, kalin: bool = False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after  = Pt(8)
    p.paragraph_format.line_spacing = Pt(17)
    run = p.add_run(metin)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = kalin


def _madde_ekle(doc: Document, metin: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(metin)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)


def _kapak_ekle(doc: Document, proje_adi: str, takim_adi: str):
    doc.add_paragraph()  # boşluk
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(proje_adi)
    r.font.name = "Arial"
    r.font.bold = True
    r.font.size = Pt(20)
    _set_color(r, KOYU_MAVI)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(6)
    r2 = p2.add_run("Proje Raporu — ProjeDanışmanAI v2")
    r2.font.name = "Arial"
    r2.font.size = Pt(13)
    _set_color(r2, ORTA_MAVI)

    doc.add_paragraph()

    meta = [
        ("Takım / Kullanıcı", takim_adi),
        ("Tarih",             datetime.now().strftime("%d.%m.%Y")),
        ("Hazırlayan",        "ProjeDanışmanAI v2 — Qwen3-14B"),
    ]
    tablo = doc.add_table(rows=len(meta), cols=2)
    tablo.style = "Table Grid"
    for i, (etiket, deger) in enumerate(meta):
        sol = tablo.rows[i].cells[0]
        sag = tablo.rows[i].cells[1]
        r1 = sol.paragraphs[0].add_run(etiket)
        r1.font.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(11)
        _set_color(r1, KOYU_MAVI)
        r2 = sag.paragraphs[0].add_run(deger)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)

    doc.add_page_break()


# ── Ana fonksiyon ─────────────────────────────────────────────────────────────

