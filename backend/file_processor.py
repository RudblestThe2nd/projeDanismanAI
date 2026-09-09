"""
file_processor.py
PDF ve DOCX dosyalarından metin çıkarır, chunk'lara böler.
ProjeDanışmanAI v2
"""

import io
import re
from typing import List, Dict


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """PDF'ten düz metin çıkarır."""
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                pages.append(f"[Sayfa {i+1}]\n{text.strip()}")
        return "\n\n".join(pages)
    except ImportError:
        raise ImportError("pypdf kurulu değil. 'pip install pypdf' çalıştır.")
    except Exception as e:
        raise ValueError(f"PDF okuma hatası: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """DOCX'ten düz metin çıkarır."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())
        return "\n\n".join(paragraphs)
    except ImportError:
        raise ImportError("python-docx kurulu değil. 'pip install python-docx' çalıştır.")
    except Exception as e:
        raise ValueError(f"DOCX okuma hatası: {str(e)}")


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Dosya türüne göre metin çıkarır."""
    fname = filename.lower()
    if fname.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif fname.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif fname.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Desteklenmeyen dosya türü: {filename}. PDF, DOCX veya TXT yükle.")


_NUMBERED_HEADING = re.compile(r"^\d+(\.\d+)*[\.\s]")


def _is_heading_para(para) -> bool:
    """Paragrafın başlık olup olmadığını belirler (stil, bold, numaralı)."""
    style_lower = para.style.name.lower()
    heading_styles = {"heading", "başlık", "title", "subtitle"}
    if any(h in style_lower for h in heading_styles):
        return True
    text = para.text.strip()
    if not text or len(text) > 120:
        return False
    # Tüm runs bold ise ve kısa ise başlık say
    runs = [r for r in para.runs if r.text.strip()]
    if runs and all(r.bold for r in runs):
        return True
    # "1." / "1.1." / "7.3." gibi numaralı kısa satırlar
    if _NUMBERED_HEADING.match(text) and len(text) < 100:
        return True
    return False


def extract_sections_from_docx(file_bytes: bytes) -> Dict[str, str]:
    """DOCX'ten {başlık: içerik} dict'i döner. Heading stili, bold ve numaralı başlıkları yakalar."""
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))

    sections: Dict[str, str] = {}
    current_heading = "__giris__"
    current_lines: List[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if _is_heading_para(para):
            if current_lines:
                sections[current_heading] = "\n".join(current_lines)
            current_heading = text
            current_lines = []
        else:
            current_lines.append(text)

    if current_lines:
        sections[current_heading] = "\n".join(current_lines)

    return sections


def extract_sections_from_pdf(file_bytes: bytes) -> Dict[str, str]:
    """PDF'ten sayfa bazlı bölümler çıkarır."""
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    sections: Dict[str, str] = {}
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            sections[f"Sayfa {i + 1}"] = text.strip()
    return sections


_TR_STOPWORDS = {
    "mi", "mı", "mu", "mü", "bir", "bu", "şu", "ve", "ile", "de", "da",
    "ki", "ya", "için", "ama", "fakat", "ancak", "var", "yok", "ne",
    "nasıl", "hangi", "neden", "sen", "ben", "biz", "siz", "değil",
    "gibi", "kadar", "daha", "en", "çok", "az", "sence", "bana", "bence",
    "acaba", "olan", "ile", "veya", "ya", "da", "bile", "misin", "musun",
    "mısın", "misiniz", "kısmını", "kısmı", "bölümünü", "bölümü",
    "değerlendirir", "değerlendirmisin", "değerlendirme", "incelemisin",
    "bakabilir", "bakarmısın", "söyler", "anlatır", "hakkında", "konusunda",
    "ilgili", "olan", "olarak", "olan", "göre", "kadar", "sonra", "önce",
}

_TR_SUFFIXES = (
    "ları", "leri", "ını", "ini", "unu", "ünü", "nın", "nin", "nun", "nün",
    "ında", "inde", "unda", "ünde", "dan", "den", "tan", "ten", "nda", "nde",
    "la", "le", "ya", "ye", "a", "e", "ı", "i", "u", "ü",
)


